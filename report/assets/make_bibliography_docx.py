"""
Generates report/bibliography.docx — the bibliography formatted in APA 7th
edition style (Times New Roman 12, double-spaced, hanging indent, italics),
ready to be opened in Word or copied directly into the memoria.

Mirrors the three groups of Section 10 (Bibliografia) of report/memoria.md.
Run from the project root:
    .venv/bin/python report/assets/make_bibliography_docx.py
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Inches, Pt

# Each entry is a list of (text, italic) runs, per APA italics rules:
# journal name + volume italic; article titles roman; thesis/report/preprint/
# webpage/dataset titles italic.
ENTRIES = {
    "1. Academic Literature": [
        [
            ("Baattite, A. (2023). ", False),
            ("Machine learning-based football tactic and style analysis", True),
            (" [Capstone design final report, School of Science and "
             "Engineering, Al Akhawayn University].", False),
        ],
        [
            ("Bolin, N., & Vesterlind, E. (2025, June). ", False),
            ("From patterns to playbooks: Identifying offensive and defensive "
             "team playing styles from possession chains with machine learning "
             "and its possible practical applications in management of football "
             "teams", True),
            (" [Master\u2019s thesis, Uppsala University, Faculty of Science and "
             "Technology]. (UPTEC I 25002)", False),
        ],
        [
            ("Gu, L., Wang, W., Plakias, S., & Zhang, S. (2025). Playing style "
             "identification in team sports: A systematic review from "
             "statistical dimensionality reduction to unsupervised machine "
             "learning. ", False),
            ("International Journal of Sports Science & Coaching", True),
            (", 1\u201320. https://doi.org/10.1177/17479541251372586", False),
        ],
        [
            ("Plakias, S., Armatas, V., Mprotsis, T., Kyranoudis, A., Gu, L., "
             "Xylomenos, F., & Michailidis, Y. (2025). What playing styles best "
             "represent the overall tactical identity of European soccer "
             "teams? ", False),
            ("Journal of Physical Education and Sport", True),
            (", ", False),
            ("25", True),
            ("(11), 2249\u20132257. https://doi.org/10.7752/jpes.2025.11248", False),
        ],
        [
            ("Tang, Z., Wang, X., & Zhang, S. (n.d.). ", False),
            ("Clustering football game situations via deep representation "
             "learning", True),
            (" [Unpublished manuscript]. (Venue, year and DOI not stated in "
             "the supplied PDF.)", False),
        ],
        [
            ("Wang, Z., Veli\u010dkovi\u0107, P., Hennes, D., Toma\u0161ev, N., "
             "Prince, L., Kaisers, M., Bachrach, Y., Elie, R., Wenliang, L. K., "
             "Piccinini, F., Spearman, W., Graham, I., Connor, J., Yang, Y., "
             "Recasens, A., Khan, M., Beauguerlange, N., Sprechmann, P., "
             "Moreno, P., ... Tuyls, K. (2023, October). ", False),
            ("TacticAI: An AI assistant for football tactics", True),
            (". arXiv. https://arxiv.org/abs/2310.10553", False),
        ],
        [
            ("Zhong, Y., Xu, Y., Zhu, K., Diaz-Cidoncha Garcia, J., "
             "G\u00f3mez Ruano, M. \u00c1., & Yi, Q. (2026). Data-driven "
             "classification of playing styles and match outcome prediction in "
             "UEFA Champions League teams. ", False),
            ("Biology of Sport", True),
            (", ", False),
            ("43", True),
            (", 575\u2013586. https://doi.org/10.5114/biolsport.2026.154944", False),
        ],
    ],
    "1.1 Cited by the Papers Above, Identified but Not Read in Full": [
        [
            ("Cintia, P., & Pappalardo, L. (2021). ", False),
            ("Coach2vec: Autoencoding the playing style of soccer coaches", True),
            (". arXiv. https://arxiv.org/abs/2106.15444", False),
        ],
        [
            ("Plakias, S., Kokkotis, C., Moustakidis, S., Tsatalas, T., "
             "Papalexi, M., Kasioura, C., Giakas, G., & Tsaopoulos, D. (2023). "
             "Identifying playing styles of European soccer teams during the "
             "key moments of the game. ", False),
            ("Journal of Physical Education and Sport", True),
            (", ", False),
            ("23", True),
            ("(4), 878\u2013890. https://doi.org/10.7752/jpes.2023.04111", False),
        ],
        [
            ("Rein, R., & Memmert, D. (2016). Big data and tactical analysis in "
             "elite soccer: Future challenges and opportunities for sports "
             "science. ", False),
            ("SpringerPlus", True),
            (", ", False),
            ("5", True),
            (", 1410. https://doi.org/10.1186/s40064-016-3108-2", False),
        ],
    ],
    "2. Manager Tenure and Tactical Philosophy Sources": [
        [
            ("Coaches\u2019 Voice. (2021). ", False),
            ("Pep Guardiola: Coach watch", True),
            (". https://www.coachesvoice.com (Retrieved August 10, 2026)", False),
        ],
        [
            ("Coaches\u2019 Voice. (2026). ", False),
            ("The tactical evolution of Pep Guardiola\u2019s Manchester City", True),
            (". https://www.coachesvoice.com (Retrieved August 10, 2026)", False),
        ],
        [
            ("Coaches\u2019 Voice. (n.d.). ", False),
            ("Roberto De Zerbi tactics and style of play", True),
            (". https://www.coachesvoice.com (Retrieved August 10, 2026)", False),
        ],
        [
            ("Opta Analyst / The Analyst. (2019\u20132025). ", False),
            ("Data-confirmed statistical pieces on Pep Guardiola\u2019s "
             "Manchester City tenure", True),
            (" [Multiple articles]. https://theanalyst.com "
             "(Retrieved August 10, 2026)", False),
        ],
        [
            ("Perarnau, M. (n.d.). ", False),
            ("Pep Confidential", True),
            (" [Serialised excerpts]. Marti Perarnau Magazine. "
             "https://www.martiperarnau.com (Retrieved August 10, 2026)", False),
        ],
        [
            ("SHfootball. (n.d.). ", False),
            ("Analysis of Roberto De Zerbi\u2019s goalkeeper and centre-back "
             "build-up mechanism", True),
            (" [Blog post]. Medium. https://medium.com "
             "(Retrieved August 10, 2026)", False),
        ],
        [
            ("SoccerTutor.com. (n.d.). ", False),
            ("Analysis of Roberto De Zerbi\u2019s build-up structure", True),
            (". https://www.soccertutor.com (Retrieved August 10, 2026)", False),
        ],
        [
            ("Spielverlagerung. (2016). ", False),
            ("Previewing Pep Guardiola\u2019s Manchester City reign", True),
            (". https://spielverlagerung.com (Retrieved August 10, 2026)", False),
        ],
        [
            ("Spielverlagerung. (2017). ", False),
            ("Manchester City: Positional play analysis", True),
            (". https://spielverlagerung.com (Retrieved August 10, 2026)", False),
        ],
        [
            ("Total Football Analysis. (2026). ", False),
            ("Tactical analysis pieces on Guardiola, Klopp and De Zerbi", True),
            (" [Multiple articles]. https://totalfootballanalysis.com "
             "(Retrieved August 10, 2026)", False),
        ],
        [
            ("Wikipedia. (2026). ", False),
            ("Managerial statistics tables: Pep Guardiola; J\u00fcrgen Klopp; "
             "Roberto De Zerbi", True),
            (". https://www.wikipedia.org (Retrieved August 10, 2026). "
             "Every tenure entry cross-corroborated against club and press "
             "sources, with every conflict explicitly resolved and documented "
             "in data/external/manager_tenures.md \u00a73.", False),
        ],
    ],
    "3. Data Source Documentation": [
        [
            ("Mathien, H. (n.d.). ", False),
            ("European soccer database", True),
            (" [Data set]. Kaggle. https://www.kaggle.com/datasets/hugomathien/"
             "soccer (Retrieved August 10, 2026). Investigated as a possible "
             "historical extension for the Bayern/Dortmund eras; found "
             "unsuitable \u2014 see decision_log.md D-005.", False),
        ],
        [
            ("soccerdata", True),
            (" [Computer software]. (n.d.). PyPI. "
             "https://pypi.org/project/soccerdata/ "
             "(Retrieved August 10, 2026). Tested for direct FBref/Understat "
             "access; insufficient on its own \u2014 see "
             "outputs/audit/data_availability_matrix.md \u00a76.", False),
        ],
        [
            ("StatsBomb. (n.d.). ", False),
            ("StatsBomb open data", True),
            (" [Data set]. GitHub. https://github.com/statsbomb/open-data "
             "(Retrieved August 10, 2026). Competition and match manifests, "
             "event-level data, match-level manager attribution; used as a "
             "separate, non-pooled evidence stream per decision_log.md D-003.", False),
        ],
        [
            ("Ziv, J. (n.d.). ", False),
            ("worldfootballR_data", True),
            (" [Data set]. GitHub. https://github.com/JaseZiv/"
             "worldfootballR_data (Retrieved August 10, 2026). Primary source: "
             "FBref/Opta match-level advanced statistics mirror, "
             "fb_advanced_match_stats release \u2014 see decision_log.md D-002.", False),
        ],
    ],
}


def add_entry(doc, runs):
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.left_indent = Inches(0.5)
    fmt.first_line_indent = Inches(-0.5)
    fmt.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    fmt.space_after = Pt(0)
    for text, italic in runs:
        r = p.add_run(text)
        r.italic = italic


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    title = doc.add_paragraph("Bibliography")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].bold = True

    for heading, entries in ENTRIES.items():
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        r = h.add_run(heading)
        r.bold = True
        for runs in entries:
            add_entry(doc, runs)

    out = "report/bibliography.docx"
    doc.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
